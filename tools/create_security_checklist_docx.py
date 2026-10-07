from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "dokumen" / "Daftar_Uji_Keamanan_dan_Ketahanan_Siber_SABABUKA.docx"


def set_cellless_font(run, name="Arial", size=Pt(10), bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = size
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    set_cellless_font(run, size=Pt(8.5), color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def keep_with_next(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    ppr.append(keep)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.07

for style_name, size, bold in [
    ("Title", 20, True),
    ("Subtitle", 11, False),
    ("Heading 1", 14, True),
    ("Heading 2", 11.5, True),
]:
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)

styles["Title"].paragraph_format.space_after = Pt(10)
styles["Title"].paragraph_format.keep_with_next = True
styles["Subtitle"].paragraph_format.space_after = Pt(18)
styles["Heading 1"].paragraph_format.space_before = Pt(14)
styles["Heading 1"].paragraph_format.space_after = Pt(7)
styles["Heading 1"].paragraph_format.keep_with_next = True
styles["Heading 2"].paragraph_format.space_before = Pt(9)
styles["Heading 2"].paragraph_format.space_after = Pt(5)
styles["Heading 2"].paragraph_format.keep_with_next = True

# Remove the decorative border inherited by Word's built-in Title style.
title_style_ppr = styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

if "Daftar Strip" not in styles:
    strip_style = styles.add_style("Daftar Strip", WD_STYLE_TYPE.PARAGRAPH)
else:
    strip_style = styles["Daftar Strip"]
strip_style.base_style = styles["Normal"]
strip_style.font.name = "Arial"
strip_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
strip_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
strip_style.font.size = Pt(10)
strip_style.paragraph_format.left_indent = Cm(0.55)
strip_style.paragraph_format.first_line_indent = Cm(-0.42)
strip_style.paragraph_format.space_after = Pt(2.5)
strip_style.paragraph_format.line_spacing = 1.04

header = section.header
hp = header.paragraphs[0]
hp.text = "SABABUKA BERSINAR  |  DAFTAR UJI KEAMANAN"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in hp.runs:
    set_cellless_font(run, size=Pt(8), bold=True, color="666666")

footer = section.footer
fp = footer.paragraphs[0]
add_page_number(fp)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Daftar Uji Keamanan dan Ketahanan Siber Sistem SABABUKA")

subtitle = doc.add_paragraph(style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Dokumen kerja untuk pelaksanaan pengujian keamanan dan penerimaan sistem")

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("SABABUKA Bersinar\nVersi 1.0  |  5 Oktober 2026")
set_cellless_font(r, size=Pt(9.5), color="444444")
meta.paragraph_format.space_after = Pt(20)

intro = doc.add_paragraph()
intro.add_run(
    "Dokumen ini memuat ruang lingkup minimum pengujian keamanan yang harus dilaksanakan sebelum "
    "SABABUKA digunakan pada lingkungan produksi. Hasil pengujian wajib dilengkapi bukti teknis, "
    "tindak lanjut perbaikan, dan pengujian ulang. Sistem dinyatakan layak apabila seluruh kriteria "
    "penerimaan pada bagian akhir dokumen telah dipenuhi."
)
intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def h1(text):
    return doc.add_paragraph(text, style="Heading 1")


def h2(text):
    return doc.add_paragraph(text, style="Heading 2")


def dash(text):
    p = doc.add_paragraph(style="Daftar Strip")
    p.add_run("- " + text)
    return p


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    properties.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    properties.append(size)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def dash_link(label, url):
    p = doc.add_paragraph(style="Daftar Strip")
    p.add_run("- ")
    add_hyperlink(p, label, url)
    return p


def dashes(items):
    for item in items:
        dash(item)


h1("1 Tujuan Pengujian")
dashes([
    "Memastikan kerahasiaan, keutuhan, ketersediaan, keaslian, akuntabilitas, dan kenirsangkalan data serta layanan SABABUKA.",
    "Memastikan setiap fungsi hanya dapat digunakan oleh pengguna, role, dan organisasi yang berwenang.",
    "Memastikan data indikator, metadata, bukti dukung, hasil verifikasi, dan publikasi tidak dapat diubah tanpa otorisasi dan catatan audit.",
    "Memastikan server, jaringan, database, aplikasi, API, Cloudflare, dan proses operasional memiliki pengamanan yang memadai.",
    "Memastikan layanan dapat dipulihkan setelah gangguan, kerusakan data, atau insiden siber.",
    "Menyediakan dasar penerimaan keamanan sebelum sistem digunakan pada lingkungan produksi.",
])

h1("2 Acuan Pengujian")
dashes([
    "Peraturan Presiden Nomor 95 Tahun 2018 tentang Sistem Pemerintahan Berbasis Elektronik.",
    "Peraturan Badan Siber dan Sandi Negara Nomor 4 Tahun 2021 tentang Pedoman Manajemen Keamanan Informasi SPBE serta Standar Teknis dan Prosedur Keamanan SPBE.",
    "Peraturan Pemerintah Nomor 71 Tahun 2019 tentang Penyelenggaraan Sistem dan Transaksi Elektronik.",
    "Undang-Undang Nomor 27 Tahun 2022 tentang Pelindungan Data Pribadi.",
    "Peraturan Presiden Nomor 39 Tahun 2019 tentang Satu Data Indonesia.",
    "Peraturan Badan Siber dan Sandi Negara Nomor 1 Tahun 2024 tentang Pengelolaan Insiden Siber.",
    "OWASP Application Security Verification Standard versi 5.0 dengan target minimum Level 2.",
    "OWASP Top 10 versi 2025.",
    "OWASP API Security Top 10 versi 2023.",
    "ISO IEC 27001:2022 sebagai acuan pengelolaan keamanan informasi organisasi.",
])

h1("3 Ruang Lingkup Aset")
dashes([
    "Domain utama, subdomain pengembangan, subdomain produksi, dan seluruh nama host yang digunakan SABABUKA.",
    "Alamat IPv4 dan IPv6 server, origin server, reverse proxy, dan layanan yang berada di belakang Cloudflare.",
    "Frontend web, backend API, database, penyimpanan file, layanan backup, monitoring, dan logging.",
    "Konfigurasi Nginx, TLS, firewall, SSH, container, sistem operasi, dan scheduler.",
    "Seluruh endpoint API yang tercantum maupun yang belum tercantum pada dokumentasi OpenAPI.",
    "Modul pengguna, organisasi, role, permission, menu, konfigurasi sistem, dan feature flag.",
    "Modul fokus kebijakan, kategori, indikator, satuan, periode, dan organisasi penanggung jawab.",
    "Modul pengajuan data OPD, verifikasi, bukti dukung, pengembalian, persetujuan, dan publikasi.",
    "Dashboard operasional, dashboard pimpinan, notifikasi, audit log, dan AI Assistant.",
    "Integrasi pihak ketiga yang digunakan sekarang atau disiapkan untuk tahap berikutnya.",
])

h1("4 Dokumen dan Persiapan Sebelum Pengujian")
dashes([
    "Daftar aset, domain, subdomain, alamat IP, port, service, dan pemilik aset tersedia dan mutakhir.",
    "Diagram arsitektur infrastruktur dan Data Flow Diagram tersedia.",
    "Dokumentasi OpenAPI sesuai dengan implementasi backend terakhir.",
    "Matriks role, permission, menu, dan batas kewenangan setiap organisasi tersedia.",
    "Klasifikasi data publik, internal, terbatas, rahasia, dan data pribadi telah ditetapkan.",
    "Daftar komponen perangkat lunak atau Software Bill of Materials tersedia.",
    "Akun pengujian tersedia untuk pengguna tanpa login, OPD, verifikator, BAPPERIDA, pimpinan, pengelola publikasi, dan superadmin.",
    "Akun nonaktif, akun tanpa permission, dan akun yang masa berlakunya berakhir disediakan untuk skenario negatif.",
    "Data uji tidak menggunakan data pribadi nyata tanpa persetujuan dan pengamanan yang sesuai.",
    "Backup telah dibuat dan diverifikasi sebelum pengujian yang berpotensi memengaruhi data.",
    "Rules of Engagement menetapkan ruang lingkup, jadwal, sumber IP penguji, larangan, kontak darurat, dan prosedur penghentian pengujian.",
    "Pengujian ketahanan dilaksanakan pada staging yang setara dengan produksi kecuali terdapat persetujuan khusus.",
])

h1("5 Pengujian Jaringan dan VPS")
h2("5.1 Paparan layanan")
dashes([
    "Lakukan pemindaian seluruh port TCP dan UDP pada IPv4 serta IPv6.",
    "Pastikan hanya port layanan yang diperlukan yang dapat diakses dari internet.",
    "Pastikan database, Docker socket, dashboard monitoring, port administrasi, dan service internal tidak dapat diakses langsung dari internet.",
    "Pastikan firewall menggunakan kebijakan default deny dan hanya memiliki rule yang terdokumentasi.",
    "Pastikan tidak terdapat perbedaan rule firewall antara IPv4 dan IPv6.",
    "Uji respons server terhadap port scanning, connection flood, dan percobaan akses berulang.",
    "Periksa service banner, informasi versi, halaman default, directory listing, dan file konfigurasi yang terekspos.",
])
h2("5.2 Cloudflare dan DNS")
dashes([
    "Uji kebocoran alamat origin melalui DNS history, subdomain, sertifikat, email header, atau konfigurasi lama.",
    "Pastikan origin hanya menerima trafik yang berasal dari jalur yang disetujui.",
    "Uji bypass Cloudflare menggunakan alamat origin, IPv6, host header, dan nama domain alternatif.",
    "Periksa dangling DNS record dan potensi subdomain takeover.",
    "Pastikan DNS record pengembangan dan produksi tidak mengarah ke layanan yang sudah tidak digunakan.",
    "Validasi konfigurasi WAF, bot protection, rate limit, dan pengecualian rule.",
])
h2("5.3 SSH dan administrasi server")
dashes([
    "Nonaktifkan login SSH langsung sebagai root.",
    "Nonaktifkan autentikasi SSH menggunakan password setelah akses SSH key terverifikasi.",
    "Gunakan akun administrasi non-root dengan sudo dan catatan audit.",
    "Batasi akses SSH berdasarkan VPN, allowlist IP, atau jalur administrasi yang disetujui.",
    "Aktifkan perlindungan brute force dan pembatasan percobaan login.",
    "Cabut SSH key, akun, dan credential yang sudah tidak digunakan.",
    "Pastikan private key, password, token, dan command history sensitif tidak tersimpan pada lokasi yang dapat diakses pihak lain.",
])
h2("5.4 Sistem operasi dan container")
dashes([
    "Pasang pembaruan keamanan sistem operasi, kernel, dan package sebelum pengujian final.",
    "Lakukan vulnerability scan pada sistem operasi dan package yang terpasang.",
    "Pastikan container berjalan sebagai pengguna non-root jika tidak ada kebutuhan teknis khusus.",
    "Batasi Linux capabilities, filesystem write access, proses, CPU, RAM, dan storage container.",
    "Gunakan image dengan versi tetap dan hindari tag latest pada produksi.",
    "Lakukan pemindaian CVE pada seluruh image container.",
    "Pastikan secret tidak tersimpan di Dockerfile, image layer, repository, atau output log.",
    "Aktifkan rotasi log dan batas ukuran untuk mencegah disk penuh.",
])

h1("6 Pengujian TLS dan HTTP Security")
dashes([
    "Pastikan seluruh akses menggunakan HTTPS dan HTTP dialihkan secara permanen ke HTTPS.",
    "Gunakan TLS versi yang masih didukung dan cipher suite yang aman.",
    "Pastikan sertifikat valid, rantai sertifikat lengkap, nama domain sesuai, dan masa berlaku dipantau.",
    "Aktifkan HTTP Strict Transport Security setelah seluruh subdomain siap menggunakan HTTPS.",
    "Aktifkan Content Security Policy dan uji seluruh halaman agar tidak bergantung pada unsafe inline atau sumber yang tidak diperlukan.",
    "Aktifkan X Content Type Options, Referrer Policy, dan Permissions Policy.",
    "Lindungi aplikasi dari clickjacking menggunakan CSP frame ancestors.",
    "Pastikan CORS hanya mengizinkan origin, method, dan header yang diperlukan.",
    "Pastikan response tidak membocorkan versi server, framework, atau detail infrastruktur.",
    "Pastikan halaman privat dan response berisi data sensitif tidak disimpan pada shared cache.",
])

h1("7 Pengujian Authentication dan Session")
dashes([
    "Uji kemungkinan enumerasi username atau email melalui login, aktivasi, dan pesan kesalahan.",
    "Uji brute force, credential stuffing, password spraying, dan percobaan login terdistribusi.",
    "Pastikan rate limit berlaku per akun dan per sumber trafik dengan respons yang tidak membocorkan status akun.",
    "Uji kebijakan panjang password, password umum, password lemah, dan penggunaan ulang password.",
    "Pastikan password disimpan menggunakan algoritma hash yang sesuai dan tidak pernah dicatat dalam log.",
    "Pastikan token aktivasi acak, memiliki masa berlaku, hanya dapat digunakan satu kali, dan tidak muncul dalam log.",
    "Uji reset password, perubahan password, dan pencabutan session setelah password berubah.",
    "Uji MFA TOTP terhadap brute force, replay, reuse, clock window, bypass, dan reset MFA.",
    "Pastikan MFA wajib untuk superadmin dan akun dengan kewenangan tinggi.",
    "Pastikan recovery code bersifat sekali pakai dan disimpan dengan aman.",
    "Uji session fixation, session hijacking, dan prediktabilitas token session.",
    "Pastikan logout, penonaktifan akun, pencabutan role, dan perubahan password membatalkan session terkait.",
    "Tetapkan idle timeout dan batas masa aktif session sesuai tingkat risiko.",
    "Pastikan cookie menggunakan Secure, HttpOnly, dan SameSite yang sesuai.",
    "Pastikan token session tidak disimpan pada localStorage atau lokasi browser yang mudah diakses script.",
    "Uji CSRF pada seluruh operasi yang mengubah data.",
    "Pastikan token, cookie, dan credential tidak muncul pada URL, referer, log, atau error response.",
])

h1("8 Pengujian Role Permission dan Isolasi OPD")
dashes([
    "Uji akses vertikal dari role OPD menuju fungsi BAPPERIDA, pimpinan, pengelola publikasi, dan superadmin.",
    "Uji akses horizontal dari OPD satu ke data, submission, bukti, dan pengguna OPD lain.",
    "Ubah organization ID, user ID, submission ID, indicator ID, evidence ID, publication ID, dan identifier lainnya pada request.",
    "Akses endpoint API secara langsung tanpa melalui menu aplikasi.",
    "Pastikan penyembunyian menu bukan satu-satunya mekanisme pembatasan akses.",
    "Uji Broken Object Level Authorization dan Broken Function Level Authorization pada seluruh endpoint.",
    "Uji mass assignment menggunakan properti JSON yang tidak ditampilkan frontend.",
    "Uji perubahan role dan permission oleh pengguna yang tidak berwenang.",
    "Uji akses akun nonaktif atau akun yang seluruh permission-nya telah dicabut.",
    "Uji akses terhadap konfigurasi sistem, feature flag, audit log, dan fungsi developer.",
    "Pastikan superadmin tidak dapat ditemukan dari endpoint publik, tetapi tetap dilindungi oleh otorisasi backend, MFA, dan audit log.",
    "Pastikan setiap penolakan akses menghasilkan status HTTP yang sesuai dan tidak membocorkan data objek.",
])

h1("9 Pengujian Alur Bisnis SABABUKA")
h2("9.1 Indikator dan metadata")
dashes([
    "Pastikan fokus kebijakan, kategori, indikator, satuan, periode, sumber data, dan OPD terkait hanya dapat diubah oleh pihak berwenang.",
    "Uji tahapan versi indikator mulai dari draft, submit, approve, activate, sampai retire.",
    "Pastikan urutan status tidak dapat dilewati melalui request langsung.",
    "Pastikan indikator aktif tidak dapat diubah tanpa pembentukan versi dan catatan audit.",
    "Pastikan kode indikator unik serta perubahan kode referensi tercatat.",
    "Pastikan metadata wajib tidak dapat dilewati melalui manipulasi API.",
    "Pastikan perubahan sumber, tahun, satuan, definisi, dan OPD penanggung jawab tercatat.",
])
h2("9.2 Pengajuan dan verifikasi data OPD")
dashes([
    "Pastikan OPD hanya dapat membuat dan mengubah pengajuan untuk organisasinya sendiri.",
    "Pastikan periode yang dikunci tidak dapat diubah melalui frontend maupun API.",
    "Pastikan data yang sudah diajukan tidak dapat diedit tanpa proses pengembalian yang sah.",
    "Pastikan verifikator hanya dapat memproses data sesuai kewenangannya.",
    "Uji pemisahan tugas antara pembuat, pemeriksa, dan pemberi persetujuan.",
    "Uji duplicate submission untuk organisasi dan periode yang sama.",
    "Uji dua pengguna yang mengubah objek yang sama secara bersamaan.",
    "Uji race condition pada submit, start review, return, approve, dan activate.",
    "Uji nilai negatif, nol, nilai sangat besar, desimal, format lokal, pembagi nol, dan nilai di luar batas yang masuk akal.",
    "Uji backdating, manipulasi waktu, replay request, dan request ganda.",
    "Pastikan operasi kritis memiliki mekanisme idempotensi atau pencegahan duplikasi yang sesuai.",
])
h2("9.3 Publikasi dan dashboard")
dashes([
    "Pastikan hanya data yang telah melewati verifikasi dan persetujuan yang dapat dipublikasikan.",
    "Pastikan data draft, dikembalikan, atau ditolak tidak tampil pada dashboard pimpinan.",
    "Pastikan paket publikasi aktif tidak dapat diubah tanpa versi atau catatan audit.",
    "Uji cache agar data pengguna atau organisasi lain tidak ditampilkan setelah pergantian akun.",
    "Pastikan sumber, tahun, satuan, status, dan waktu pembaruan tetap melekat pada data yang ditampilkan.",
    "Pastikan ekspor data mengikuti klasifikasi data dan hak akses pengguna.",
])

h1("10 Pengujian Upload dan Download Bukti Dukung")
dashes([
    "Uji ekstensi palsu, double extension, MIME type palsu, dan ketidaksesuaian magic bytes.",
    "Uji file polyglot, SVG aktif, HTML, PDF dengan script, macro Office, dan file executable.",
    "Uji path traversal, karakter khusus, Unicode, nama file sangat panjang, dan filename injection.",
    "Uji file yang melebihi batas, decompression bomb, dan unggahan simultan.",
    "Lakukan pengujian antivirus menggunakan file uji yang aman sesuai prosedur penguji.",
    "Pastikan file disimpan di luar web root dan tidak dapat dieksekusi oleh server.",
    "Pastikan nama objek pada storage tidak mudah ditebak dan tidak menggunakan nama asli sebagai satu-satunya identifier.",
    "Pastikan setiap download menjalankan pemeriksaan authorization.",
    "Pastikan URL sementara memiliki masa berlaku dan tidak dapat digunakan setelah dicabut.",
    "Pastikan file yang telah dihapus tidak lagi dapat diakses melalui URL lama, cache, atau direct object reference.",
    "Pastikan retensi, pemusnahan, karantina, dan pencatatan aktivitas file telah ditetapkan.",
])

h1("11 Pengujian Web API Input dan Database")
h2("11.1 Kerentanan aplikasi")
dashes([
    "Uji SQL injection pada seluruh parameter path, query, header, cookie, dan body.",
    "Uji stored XSS, reflected XSS, dan DOM based XSS.",
    "Uji command injection, template injection, path traversal, dan file inclusion.",
    "Uji Server Side Request Forgery pada fitur yang menerima URL atau memanggil layanan lain.",
    "Uji prototype pollution, parameter pollution, Unicode normalization, dan JSON parser abuse.",
    "Uji host header injection, open redirect, header injection, dan HTTP request smuggling.",
    "Uji CSV atau formula injection pada fitur ekspor.",
    "Uji error handling agar stack trace, query, path internal, dan secret tidak tampil pada response.",
    "Cari file env, backup, source map, repository Git, konfigurasi, dan endpoint debug yang terekspos.",
])
h2("11.2 Keamanan API")
dashes([
    "Verifikasi seluruh endpoint terdaftar dalam inventaris dan memiliki versi yang jelas.",
    "Uji Broken Object Level Authorization pada setiap endpoint yang menerima identifier objek.",
    "Uji Broken Object Property Level Authorization dan mass assignment.",
    "Uji Broken Function Level Authorization pada endpoint administrasi dan workflow.",
    "Uji pembatasan pagination, filter, body, batch, upload, dan konsumsi resource.",
    "Uji akses berlebihan terhadap alur bisnis sensitif.",
    "Uji SSRF, konfigurasi API, endpoint lama, dan konsumsi API pihak ketiga.",
    "Pastikan validasi dilakukan pada backend dan tidak bergantung pada validasi frontend.",
])
h2("11.3 Database dan data")
dashes([
    "Pastikan koneksi database hanya berasal dari backend atau jalur administrasi yang disetujui.",
    "Gunakan akun database dengan hak minimum dan pisahkan akun administrasi dari akun aplikasi.",
    "Pastikan query menggunakan parameter dan tidak merangkai input pengguna secara langsung.",
    "Pastikan transaksi menjaga konsistensi pada workflow yang terdiri dari beberapa perubahan data.",
    "Uji isolasi data antarorganisasi pada tingkat service dan query database.",
    "Pastikan data sensitif, credential, token, dan MFA secret tidak muncul pada query log.",
    "Pastikan koneksi database terenkripsi ketika melewati jaringan yang tidak sepenuhnya tepercaya.",
    "Pastikan backup database terenkripsi dan aksesnya dicatat.",
])

h1("12 Pengujian Frontend dan Browser")
dashes([
    "Pastikan bundle frontend tidak memuat password, API key, token, alamat internal, atau secret lainnya.",
    "Pastikan source map produksi tidak dapat diunduh tanpa kebutuhan yang sah.",
    "Pastikan data sensitif tidak disimpan pada localStorage, sessionStorage, IndexedDB, atau browser cache tanpa kebutuhan.",
    "Uji tombol kembali browser setelah logout agar tidak menampilkan data pengguna sebelumnya.",
    "Uji pergantian akun pada browser yang sama agar state pengguna lama tidak terbawa.",
    "Pastikan pesan error frontend tidak membocorkan detail implementasi backend.",
    "Lakukan pemindaian dependency frontend dan pastikan tidak ada kerentanan kritis atau tinggi yang belum ditangani.",
    "Uji tampilan keamanan dan fungsi utama pada browser serta ukuran layar yang didukung.",
])

h1("13 Pengujian AI Assistant")
dashes([
    "Uji prompt injection dari pertanyaan pengguna, metadata, data indikator, dan dokumen sumber.",
    "Uji permintaan untuk mengabaikan role, permission, instruksi sistem, dan batas organisasi.",
    "Pastikan AI hanya dapat menggunakan data yang boleh dilihat oleh pengguna yang sedang login.",
    "Uji kebocoran data OPD lain, prompt sistem, token, credential, dan konfigurasi internal.",
    "Pastikan jawaban menyertakan rujukan data yang benar dan tidak mengarang sumber.",
    "Pastikan jawaban AI tidak dianggap sebagai keputusan administratif final.",
    "Pastikan AI tidak dapat mengubah data atau menjalankan fungsi administratif tanpa otorisasi terpisah.",
    "Batasi panjang prompt, jumlah permintaan, waktu proses, dan konsumsi biaya.",
    "Uji konten berbahaya yang tersimpan di sumber data atau dokumen.",
    "Pastikan data pribadi tidak dikirim kepada penyedia eksternal tanpa dasar, perjanjian, dan perlindungan yang sesuai.",
    "Tetapkan retensi serta kontrol akses untuk prompt dan respons yang disimpan.",
    "Pastikan kegagalan layanan AI tidak mengganggu fungsi utama SABABUKA.",
])

h1("14 Logging Audit dan Monitoring")
dashes([
    "Catat login berhasil, login gagal, logout, aktivasi akun, perubahan password, perubahan MFA, dan penggunaan recovery code.",
    "Catat penambahan atau pencabutan role, permission, menu, akun, dan organisasi.",
    "Catat perubahan konfigurasi sistem dan feature flag.",
    "Catat pembuatan serta perubahan fokus kebijakan, kategori, indikator, periode, dan metadata.",
    "Catat submit, start review, return, approve, activation, publication, dan perubahan status lainnya.",
    "Catat upload, download, dan penghapusan bukti dukung.",
    "Catat penggunaan fungsi superadmin dan percobaan akses yang ditolak.",
    "Simpan actor, waktu, alamat sumber, objek, tindakan, hasil, dan perubahan nilai yang relevan.",
    "Pastikan pengguna aplikasi tidak dapat mengubah atau menghapus audit log.",
    "Pastikan waktu server tersinkronisasi dan menggunakan zona waktu yang konsisten.",
    "Pastikan log tidak memuat password, TOTP secret, cookie, token, atau data pribadi yang tidak diperlukan.",
    "Aktifkan alert untuk brute force, privilege escalation, perubahan konfigurasi kritis, dan unduhan massal.",
    "Kirim log penting ke lokasi terpisah atau terpusat agar tetap tersedia ketika server utama terganggu.",
    "Tetapkan retensi log, akses petugas, dan prosedur penelusuran insiden.",
])

h1("15 Uji Ketahanan dan Ketersediaan")
dashes([
    "Tetapkan jumlah pengguna bersamaan, response time, throughput, dan error rate yang menjadi target layanan.",
    "Lakukan load test untuk login, dashboard, daftar indikator, pengajuan, pencarian, dan publikasi.",
    "Lakukan stress test sampai batas aman untuk mengetahui titik jenuh dan pola kegagalan.",
    "Uji lonjakan trafik, unggahan simultan, query berat, pagination besar, dan request body besar.",
    "Uji slow request, connection exhaustion, database connection exhaustion, dan storage penuh.",
    "Uji restart frontend, backend, container, database, dan reverse proxy.",
    "Uji gangguan koneksi database dan layanan eksternal.",
    "Pastikan rate limit melindungi layanan tanpa memblokir pengguna sah secara permanen.",
    "Pastikan kegagalan satu komponen tidak langsung menyebabkan kegagalan berantai.",
    "Pastikan aplikasi pulih otomatis atau memiliki prosedur pemulihan yang telah diuji.",
    "Pengujian DDoS hanya dilakukan dengan persetujuan dan koordinasi penyedia VPS serta Cloudflare.",
])

h1("16 Backup Disaster Recovery dan Insiden")
h2("16.1 Backup dan pemulihan")
dashes([
    "Backup mencakup database, file bukti, konfigurasi penting, dan data yang diperlukan untuk pemulihan.",
    "Backup dilakukan otomatis, dipantau, dienkripsi, dan memiliki salinan di lokasi berbeda.",
    "Sediakan salinan backup yang tidak dapat diubah langsung dari server aplikasi.",
    "Tetapkan retensi backup harian, mingguan, dan bulanan.",
    "Lakukan pemeriksaan integritas backup secara berkala.",
    "Lakukan restore drill ke lingkungan baru dan verifikasi konsistensi data hasil pemulihan.",
    "Tetapkan Recovery Point Objective dan Recovery Time Objective.",
    "Catat hasil restore drill, waktu pemulihan, masalah, dan tindakan koreksi.",
])
h2("16.2 Penanganan insiden")
dashes([
    "Tersedia daftar kontak tim tanggap insiden, Diskominfo, pengelola sistem, pimpinan, dan penyedia layanan.",
    "Tersedia prosedur deteksi, triase, isolasi, preservasi bukti, eradikasi, pemulihan, dan pelaporan.",
    "Lakukan simulasi pengambilalihan akun superadmin.",
    "Lakukan simulasi kebocoran data pribadi dan akses lintas OPD.",
    "Lakukan simulasi kerusakan database, ransomware, dan tidak tersedianya server.",
    "Lakukan tabletop exercise sebelum go live dan secara berkala setelah sistem beroperasi.",
    "Pastikan hasil simulasi menghasilkan daftar perbaikan, PIC, dan tenggat penyelesaian.",
])

h1("17 Secure Development dan Supply Chain")
dashes([
    "Lakukan Static Application Security Testing pada source code backend dan frontend.",
    "Lakukan manual secure code review pada authentication, authorization, workflow, upload, dan publikasi.",
    "Lakukan Software Composition Analysis pada dependency backend dan frontend.",
    "Lakukan secret scanning pada repository dan riwayat Git.",
    "Lakukan scan image container dan konfigurasi deployment.",
    "Lakukan Dynamic Application Security Testing pada staging.",
    "Lakukan fuzzing pada parser dan endpoint penting dengan batas yang disetujui.",
    "Gunakan dependency lockfile dan registry yang tepercaya.",
    "Ganti dependency yang tidak lagi dipelihara atau memiliki kerentanan yang tidak dapat dimitigasi.",
    "Terapkan branch protection, peer review, dan persetujuan perubahan untuk kode produksi.",
    "Pisahkan secret development, staging, dan production.",
    "Rotasi secret setelah pentest dan sebelum go live.",
    "Hentikan proses deployment apabila ditemukan kerentanan kritis atau tinggi yang belum diselesaikan.",
])

h1("18 Metode Pelaksanaan Pentest")
dashes([
    "Laksanakan external reconnaissance terhadap domain, IP, DNS, port, dan paparan informasi.",
    "Laksanakan black box test tanpa akun untuk menilai permukaan serangan publik.",
    "Laksanakan authenticated grey box test menggunakan setiap role yang tersedia.",
    "Laksanakan pengujian horizontal antarorganisasi dan vertikal antarrole.",
    "Laksanakan pengujian seluruh API berdasarkan OpenAPI serta endpoint yang ditemukan selama pengujian.",
    "Laksanakan pengujian server, container, database, Cloudflare, dan konfigurasi jaringan.",
    "Laksanakan code review, dependency scan, secret scan, dan image scan.",
    "Laksanakan uji ketahanan pada lingkungan yang disetujui.",
    "Laksanakan retest setelah perbaikan menggunakan langkah reproduksi yang sama dan pengujian regresi terkait.",
])

h1("19 Isi Minimum Laporan Temuan")
dashes([
    "Nomor dan judul temuan.",
    "Aset, endpoint, parameter, dan role yang digunakan.",
    "Tanggal serta lingkungan pengujian.",
    "Tingkat risiko dan skor CVSS yang disertai penilaian dampak bisnis.",
    "Uraian kelemahan dan akar penyebab.",
    "Langkah reproduksi yang dapat diulangi.",
    "Bukti request, response, screenshot, log, atau output alat yang telah disanitasi.",
    "Data, organisasi, dan proses bisnis yang terdampak.",
    "Rekomendasi perbaikan yang spesifik.",
    "PIC, target penyelesaian, dan status tindak lanjut.",
    "Hasil retest dan bukti bahwa perbaikan efektif.",
])

h1("20 Dokumen Hasil yang Wajib Diserahkan")
dashes([
    "Surat atau dokumen ruang lingkup dan Rules of Engagement yang telah disetujui.",
    "Daftar aset dan endpoint yang benar-benar diuji.",
    "Laporan eksekutif untuk pimpinan.",
    "Laporan teknis lengkap untuk pengelola dan pengembang.",
    "Daftar temuan beserta severity, bukti, rekomendasi, PIC, dan target perbaikan.",
    "Laporan hasil pemindaian SAST, DAST, SCA, secret, container, dan infrastruktur.",
    "Software Bill of Materials versi yang diuji.",
    "Matriks hasil verifikasi OWASP ASVS Level 2.",
    "Matriks pengujian role dan isolasi antar-OPD.",
    "Hasil load test, stress test, dan uji pemulihan.",
    "Berita acara restore drill dan simulasi insiden.",
    "Laporan perbaikan dan retest.",
    "Daftar risiko yang diterima secara tertulis beserta pejabat pemberi persetujuan.",
    "Berita acara hasil akhir pengujian keamanan.",
])

h1("21 Kriteria Kelulusan")
dashes([
    "Seluruh aset, domain, API, dan role dalam ruang lingkup telah diuji dan tercantum pada laporan.",
    "Seluruh persyaratan OWASP ASVS 5.0 Level 2 yang berlaku memiliki bukti verifikasi.",
    "Tidak ada temuan Critical yang masih terbuka.",
    "Tidak ada temuan High yang masih terbuka.",
    "Tidak ada temuan Medium terbuka yang berkaitan dengan authentication, authorization, data pribadi, upload, publikasi, database, atau isolasi antar-OPD.",
    "Temuan Medium lainnya telah diperbaiki atau memiliki penerimaan risiko tertulis dari pejabat berwenang.",
    "Seluruh temuan Low dan Informational memiliki rencana tindak lanjut serta tenggat.",
    "Retest menyatakan perbaikan efektif dan tidak menimbulkan kerentanan baru.",
    "Tidak terdapat default credential, password lemah, secret, private key, atau token yang terekspos.",
    "MFA aktif untuk superadmin dan akun dengan kewenangan tinggi.",
    "RBAC backend dan isolasi antar-OPD telah dibuktikan melalui pengujian positif dan negatif.",
    "Firewall, SSH hardening, TLS, HTTP security header, WAF, dan pembatasan origin telah diverifikasi.",
    "Database dan service internal tidak dapat diakses langsung dari internet.",
    "Backup berhasil dipulihkan dan memenuhi RPO serta RTO yang ditetapkan.",
    "Monitoring, alert, audit log, dan penanganan insiden telah diuji.",
    "Dokumentasi arsitektur, OpenAPI, matriks permission, SOP, dan konfigurasi sesuai kondisi sistem yang diuji.",
    "Tidak terdapat perubahan signifikan setelah pentest. Perubahan signifikan wajib menjalani pengujian keamanan ulang.",
    "Laporan akhir dan berita acara retest telah ditandatangani oleh pihak yang berwenang.",
])

h1("22 Catatan Penerapan")
dashes([
    "Kelulusan pengujian berlaku untuk versi aplikasi, konfigurasi, aset, dan periode pengujian yang tercantum dalam laporan.",
    "Pentest diulang sebelum go live, setelah perubahan besar, setelah perubahan arsitektur atau kontrol akses, dan setelah insiden keamanan.",
    "Pemindaian otomatis tidak menggantikan pengujian manual terhadap authorization, alur bisnis, isolasi OPD, dan fungsi superadmin.",
    "Temuan dinilai berdasarkan dampak teknis dan dampak terhadap layanan pemerintahan, bukan hanya skor alat pemindai.",
    "Keputusan go live diberikan setelah seluruh syarat kelulusan dipenuhi dan risiko tersisa diterima secara tertulis.",
])

h1("23 Status Kondisi Awal yang Perlu Ditindaklanjuti")
dashes([
    "Firewall VPS terakhir terpantau belum aktif dan harus dikonfigurasi sebelum lingkungan produksi dibuka.",
    "Akses SSH langsung sebagai root menggunakan password masih tersedia dan harus diganti menjadi SSH key melalui akun administrasi non-root.",
    "Content Security Policy pada konfigurasi aplikasi masih dinonaktifkan dan perlu dirancang, diaktifkan, serta diuji sebelum produksi.",
    "Pentest independen, restore drill, dan simulasi insiden belum menjadi dasar kelulusan sampai bukti hasil pengujiannya tersedia.",
])

h1("24 Referensi")
dash_link("Peraturan Presiden Nomor 95 Tahun 2018 tentang SPBE", "https://peraturan.bpk.go.id/Details/96913/perpres-no-95-tahun-2018")
dash_link("Peraturan BSSN Nomor 4 Tahun 2021", "https://peraturan.bpk.go.id/Home/Details/174275/peraturan-bssn-no-4-tahun-2021")
dash_link("Peraturan Pemerintah Nomor 71 Tahun 2019", "https://peraturan.bpk.go.id/Details/122030/pp-no-71-tahun-2019")
dash_link("Undang-Undang Nomor 27 Tahun 2022", "https://peraturan.bpk.go.id/Home/Download/224884/UU%20Nomor%2027%20Tahun%202022.pdf")
dash_link("Peraturan Presiden Nomor 39 Tahun 2019", "https://peraturan.bpk.go.id/Details/108813/perpres-no-39-tahun-201")
dash_link("Peraturan BSSN Nomor 1 Tahun 2024", "https://peraturan.bpk.go.id/Details/291240/peraturan-bssn-no-1-tahun-2024")
dash_link("OWASP Application Security Verification Standard", "https://owasp.org/projects/asvs")
dash_link("OWASP Top 10", "https://top10.owasp.org/2025/")
dash_link("OWASP API Security", "https://api-security.owasp.org/")
dash_link("ISO IEC 27001:2022", "https://www.iso.org/standard/27001")

# Prevent a heading from being stranded at the bottom of a page.
for paragraph in doc.paragraphs:
    if paragraph.style.name in {"Title", "Heading 1", "Heading 2"}:
        keep_with_next(paragraph)

props = doc.core_properties
props.title = "Daftar Uji Keamanan dan Ketahanan Siber Sistem SABABUKA"
props.subject = "Checklist pengujian keamanan dan penerimaan sistem"
props.author = "Tim SABABUKA Bersinar"
props.keywords = "SABABUKA, keamanan, pentest, SPBE, BSSN, OWASP"
props.comments = "Dokumen kerja pengujian keamanan"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
